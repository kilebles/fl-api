"""Voice API over fl-voice_translator.

By default the translator is the stand-in in tests/fake_translator (same protocol and files). To check the real
contract, point FL_TRANSLATOR_REPO at a checkout of fl-voice_translator: the tests then run against a copy of it,
and FL_TRANSLATOR_PYTHON at a Python with its requirements for the end-to-end test with the real workers.
The bot is not started: tests change queue/*.json the way the bot does.
"""

import json
import os
import shutil
import subprocess
import sys
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from src.api.app import create_app
from src.core.settings import Settings

FAKE = Path(__file__).parent / "fake_translator"
EN_VOICE, RU_VOICE, ES_VOICE = "98762dc8579647fcbed3fd2cdeca80a1", "1c0fbdf66f484cde9bb6fd27d7012a98", "e4cd5ab1a4c64b0a9ec9bf156bc8e0bd"


def _translator(tmp_path: Path) -> Path:
    root = tmp_path / "translator"
    real = os.environ.get("FL_TRANSLATOR_REPO")
    skip = shutil.ignore_patterns(".git", "work", "output", "queue", "input", "tests", ".venv", "__pycache__",
                                  "config.local.json", "voices.json", "bot_config.json")
    shutil.copytree(real or FAKE, root, ignore=skip)
    for d in ("work", "output", "queue", "input"):
        (root / d).mkdir(exist_ok=True)
    voices = {
        "voices": [
            {"id": EN_VOICE, "name": "Голос 1 (Brian)"},
            {"id": RU_VOICE, "name": "RU 2"},
            {"id": ES_VOICE, "name": "ESP 2"},
        ],
        "profiles": [
            {
                "id": "p1",
                "name": "Голос 1 (Indians)",
                "voices": {"EN": EN_VOICE, "RU": RU_VOICE, "ES": ES_VOICE},
                "percents": {"RU": [70, 75], "ES": [90, 95]},
            }
        ],
        "active": "p1",
    }
    (root / "tools" / "voices.json").write_text(json.dumps(voices, ensure_ascii=False), encoding="utf-8")
    return root


def _settings(tmp_path: Path, voice_dir: Path) -> Settings:
    return Settings(
        _env_file=None,
        database_url=f"sqlite+aiosqlite:///{tmp_path / 'db.sqlite'}",
        work_dir=tmp_path / "work",
        storage_dir=tmp_path / "storage",
        accounts_dir=tmp_path / "accounts",
        voice_dir=voice_dir,
        voice_python=sys.executable,
        fish_api_key="",
    )


def _docx(paragraphs: list[str]) -> bytes:
    doc = Document()
    doc.add_paragraph("Introduction")
    for i, p in enumerate(paragraphs, 1):
        doc.add_paragraph(f"{i}.\t{p}")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


EN = [
    "[calm, intimate documentary narrator] How do you turn a frozen plain into a warm home? "
    "[pause] They had <emphasis>no wood for a fire</emphasis>.",
    "In 1923 the Arctic travelers built an igloo in forty-five minutes.",
]
RU = [
    "Зимой тысяча восемьсот двенадцатого года армия отступала по разорённой дороге, и каждый день уносил сотни жизней.",
    "Генерал понимал, что спасти людей может только быстрый марш к реке, где ещё стояли мосты.",
]
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class Bot:
    """Does to queue/*.json what the translator bot does."""

    def __init__(self, root: Path):
        self.root = root

    def job(self, n: int) -> dict:
        return json.loads((self.root / "queue" / f"{n:04d}.json").read_text(encoding="utf-8"))

    def save(self, job: dict) -> None:
        (self.root / "queue" / f"{job['id']:04d}.json").write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")

    def update(self, n: int, **fields) -> dict:
        job = self.job(n)
        job.update(fields)
        self.save(job)
        return job

    def add(self, **job) -> dict:
        n = max(int(f.name[:4]) for f in (self.root / "queue").glob("[0-9][0-9][0-9][0-9].json")) + 1
        job["id"] = n
        self.save(job)
        return job

    def translate(self, slug: str, lang: str, texts: list[str]) -> None:
        """run.py result: work/<slug>/<LANG>.md (+ stats and docx, via the translator itself)."""
        for i, t in enumerate(texts, 1):
            r = subprocess.run([sys.executable, "tools/api.py", "set-paragraph", slug, lang, str(i)], input=t, text=True,
                               capture_output=True, cwd=self.root, encoding="utf-8")
            assert r.returncode == 0, r.stdout + r.stderr

    def tone(self, path: Path, seconds: float) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"\xff\xfb" + b"\x00" * int(24000 * seconds))  # 192 kbps-sized stand-in


@pytest.fixture()
def env(tmp_path):
    root = _translator(tmp_path)
    with TestClient(create_app(_settings(tmp_path, root))) as client:
        yield client, Bot(root), client.app.state.voice


def _upload(client, paragraphs, name="arctic_story.docx"):
    r = client.post("/api/voice/scripts", files={"file": (name, _docx(paragraphs), DOCX)})
    assert r.status_code == 200, r.text
    return r.json()


def test_models_languages_profiles_and_norms_come_from_translator(env):
    client, bot, _ = env
    models = client.get("/api/voice/models").json()
    assert [m["id"] for m in models["voice_models"]] == ["s2.1", "drama-3"]
    langs = {lang["code"]: lang for lang in client.get("/api/voice/languages").json()}
    assert set(langs) == {"EN", "RU", "ES", "PT", "DE", "IT", "PL", "FR"}
    profs = client.get("/api/voice/profiles").json()
    assert profs[0]["norms"]["RU"] == [70, 75]  # the profile's own
    cfg = json.loads((bot.root / "tools" / "config.json").read_text(encoding="utf-8"))["languages"]
    assert profs[0]["norms"]["DE"] == [cfg["DE"]["min"], cfg["DE"]["max"]]  # the translator's config, not a copy
    assert profs[0]["norms"]["EN"] == [100, 100]


def test_translate_job_goes_to_translator_queue_and_follows_the_bot(env):
    client, bot, svc = env
    script = _upload(client, EN)
    assert script["paragraphs_total"] == 2 and script["source_lang"] == "EN" and script["emotion_tags_count"] == 1
    created = client.post("/api/voice/jobs", json={
        "script_id": script["script_id"], "mode": "translate", "targets": ["RU"], "profile_id": "p1",
        "model": "gemini-3.8-flash-medium", "voice_model": "drama-3", "carry_mood": True, "project_id": "pr1",
    }).json()
    job_id = created["id"]
    assert created["status"] == "queued" and created["model"] == "gemini-3.8-flash-medium"
    assert created["voice_model"] == "drama-3" and created["carry_mood"] is True and created["project_id"] == "pr1"
    q = bot.job(job_id)
    assert q["source"] == "web" and q["kind"] == "translate" and q["langs"] == ["RU"]
    assert q["engine"] == "agy" and q["agy_model"] == "gemini-3.8-flash-medium" and q["pipeline"] is False
    assert client.get("/api/voice/queue").json()["worker_online"] is False

    svc.scan()
    bot.update(job_id, status="running", progress={"stage": "translate", "lang": "RU", "part": 1, "parts": 1})
    (bot.root / "queue" / "bot.alive").write_text("{}")
    job = client.get(f"/api/voice/jobs/{job_id}").json()
    ru = next(lang for lang in job["languages"] if lang["lang"] == "RU")
    assert job["status"] == "running" and ru["translate"]["state"] == "running" and ru["voice"]["state"] == "skipped"
    assert client.get("/api/voice/queue").json()["worker_online"] is True
    assert client.get("/api/voice/queue").json()["slots"]["gemini"]["used"] == 1

    bot.update(job_id, pause_until="24.10 18:30", pause_reason="limit")
    paused = client.get(f"/api/voice/jobs/{job_id}").json()
    assert paused["status"] == "paused" and paused["pause_until"].startswith(f"{paused['pause_until'][:4]}-10-24T18:30")

    bot.translate(script["script_id"], "RU", [
        "[calm, intimate documentary narrator] Как превратить ледяную равнину в тёплый дом? [pause] У них <emphasis>не было дров</emphasis>.",
        "В тысяча девятьсот двадцать третьем году путники построили иглу за сорок пять минут.",
    ])
    bot.update(job_id, status="done", results={"RU": "71%"}, finished="24.10 19:00", pause_until=None)
    svc.scan()
    events = [e["event"] for e in svc._event_history]
    assert "job.done" in events
    job = client.get(f"/api/voice/jobs/{job_id}").json()
    ru = next(lang for lang in job["languages"] if lang["lang"] == "RU")
    assert job["status"] == "done" and ru["translate"]["state"] == "done" and ru["translate"]["length_pct"]
    assert [f["kind"] for f in job["files"]] == ["docx"]  # only what exists: no audio yet
    r = client.get(f"/api/voice/files/{job['files'][0]['id']}")
    assert r.status_code == 200

    # paragraphs: lengths from the translator's stats, carried mood, edits with versions
    paras = client.get(f"/api/voice/jobs/{job_id}/languages/RU/paragraphs").json()
    # profile range 70–75%; short paragraphs get the translator's extra tolerance
    assert len(paras) == 2 and paras[0]["norm"][0] <= 70 and paras[0]["norm"][1] >= 75 and paras[0]["audio"] is None
    assert paras[1]["carried_mood"] == "[calm, intimate documentary narrator]"
    p1 = client.patch(f"/api/voice/jobs/{job_id}/languages/RU/paragraphs/1",
                      json={"text": "[calm, intimate documentary narrator] Как сделать из равнины дом? [pause] "
                                    "У них <emphasis>не было дров</emphasis>.", "version": paras[0]["version"]}).json()
    assert p1["edited"] is True and p1["tags_ok"] is True and p1["version"] == paras[0]["version"] + 1
    conflict = client.patch(f"/api/voice/jobs/{job_id}/languages/RU/paragraphs/1",
                            json={"text": "Конфликт", "version": paras[0]["version"]})
    assert conflict.status_code == 409 and conflict.json()["error"]["code"] == "version_conflict"
    client.patch(f"/api/voice/jobs/{job_id}/languages/RU/paragraphs/1",
                 json={"text": "[спокойно] Как превратить пустошь в дом?", "version": p1["version"]})
    tags = client.get(f"/api/voice/jobs/{job_id}/languages/RU/paragraphs?filter=tags").json()
    assert [p["n"] for p in tags] == [1] and tags[0]["tags_ok"] is False
    reverted = client.post(f"/api/voice/jobs/{job_id}/languages/RU/paragraphs/1/revert").json()
    assert reverted["edited"] is False and reverted["text"].startswith("[calm, intimate")
    m = client.post(f"/api/voice/jobs/{job_id}/languages/RU/paragraphs/2/length", json={"text": "Коротко."}).json()
    assert m["in_norm"] is False and m["chars_delta"] < 0
    assert client.get(f"/api/voice/jobs/{job_id}/languages/RU/paragraphs/1/audio").status_code == 404


def test_all_mode_chain_voices_mixes_and_serves_real_audio(env):
    client, bot, svc = env
    script = _upload(client, EN)
    slug = script["script_id"]
    job_id = client.post("/api/voice/jobs", json={
        "script_id": slug, "mode": "all", "targets": ["RU", "ES"], "profile_id": "p1", "voices": {"ES": "v1"},
        "model": "claude-code", "voice_model": "s2.1",
    }).json()["id"]
    q = bot.job(job_id)
    assert q["pipeline"] is True and q["engine"] == "claude" and "agy_model" not in q
    assert q["voices"]["ES"] == EN_VOICE and q["voices"]["EN"] == EN_VOICE and q["voices"]["RU"] == RU_VOICE

    bot.translate(slug, "RU", ["Как превратить равнину в дом?", "В тысяча девятьсот двадцать третьем году."])
    bot.translate(slug, "ES", ["¿Cómo convertir una llanura en un hogar?", "En mil novecientos veintitrés."])
    bot.update(job_id, status="done", results={"RU": "80%", "ES": "90%"})
    common = {k: q[k] for k in ("source", "web", "voices", "voice_model", "carry_mood", "profile")}
    voice = bot.add(kind="voice", name="arctic_story.docx", status="running", pipeline=True, root=job_id,
                    parent=job_id, results={}, files=[],
                    items=[{"docx": "input/arctic_story.docx", "lang": "EN", "voice": EN_VOICE, "source": True},
                           {"docx": f"output/{slug}_RU.docx", "lang": "RU", "voice": RU_VOICE},
                           {"docx": f"output/{slug}_ES.docx", "lang": "ES", "voice": EN_VOICE}],
                    progress={"stage": "voice", "lang": "RU", "done": 1, "total": 2}, **common)
    jobs = client.get("/api/voice/jobs").json()
    assert [j["id"] for j in jobs] == [job_id]  # one web job for the whole chain
    job = jobs[0]
    ru = next(lang for lang in job["languages"] if lang["lang"] == "RU")
    assert job["status"] == "running" and ru["stage"] == "voice" and ru["voice"]["state"] == "running"
    assert ru["voice"]["done"] == 1 and ru["translate"]["state"] == "done"

    audio = bot.root / "output" / "audio"
    for n in (1, 2):
        bot.tone(audio / slug / f"{n}.mp3", 2.0)
        bot.tone(audio / f"{slug}_RU" / f"{n}.mp3", 1.8)
    bot.update(voice["id"], status="done", results={"EN": "2 дорожек", "RU": "2 дорожек", "ES": "голос недоступен"},
               skipped={"ES": EN_VOICE})
    mix = bot.add(kind="mix", name=slug, status="pending", script=slug, pipeline=True, root=job_id, parent=voice["id"],
                  results={}, files=[], langs=[], source="web", web=q["web"])
    svc.scan()
    job = client.get(f"/api/voice/jobs/{job_id}").json()
    es = next(lang for lang in job["languages"] if lang["lang"] == "ES")
    assert es["voice"]["state"] == "voice_missing" and es["mix"]["state"] == "skipped"
    assert job["summary"]["skipped"]["ES"]["code"] == "voice_missing"

    for n in (1, 2):
        bot.tone(bot.root / "output" / "mix" / slug / "EN" / f"{n}.mp3", 2.5)
        bot.tone(bot.root / "output" / "mix" / slug / "RU" / f"{n}.mp3", 2.5)
    bot.update(mix["id"], status="done", langs=["EN", "RU"], finished="24.10 20:00",
               mix={"master": "EN", "duration": 5.0, "paragraphs": 2, "extended": 0,
                    "stats": {"EN": {"sped": 0, "padded": 0, "max_speed": 1.0},
                              "RU": {"sped": 1, "padded": 1, "max_speed": 1.04}}})
    svc.scan()
    assert any(e["event"] == "job.done" and e["data"]["job_id"] == job_id for e in svc._event_history)
    job = client.get(f"/api/voice/jobs/{job_id}").json()
    ru = next(lang for lang in job["languages"] if lang["lang"] == "RU")
    assert job["status"] == "done" and ru["mix"]["state"] == "done" and ru["mix"]["sped_up"] == 1
    assert job["source_duration_sec"] == 5.0
    kinds = {(f["lang"], f["kind"]) for f in job["files"]}
    assert {("RU", "docx"), ("RU", "raw_zip"), ("RU", "mixed_zip"), ("EN", "raw_zip")} <= kinds
    assert ("ES", "raw_zip") not in kinds

    paras = client.get(f"/api/voice/jobs/{job_id}/languages/RU/paragraphs").json()
    assert paras[0]["audio"]["source_sec"] == pytest.approx(2.0, abs=0.05)
    assert paras[0]["audio"]["slot_sec"] == pytest.approx(2.5, abs=0.05)
    assert client.get(f"/api/voice/jobs/{job_id}/languages/RU/paragraphs/1/audio").status_code == 200
    assert client.get(f"/api/voice/jobs/{job_id}/languages/RU/paragraphs/1/audio?track=source").status_code == 200
    mixed = next(f for f in job["files"] if f["lang"] == "RU" and f["kind"] == "mixed_zip")
    r = client.get(mixed["url"])
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert client.get(f"/api/voice/jobs/{job_id}/zip").status_code == 200

    # missing voice: revoice the skipped language with another voice → a new step in the same chain
    re_job = client.post(f"/api/voice/jobs/{job_id}/revoice-skipped", json={"voices": {"ES": "v3"}})
    assert re_job.status_code == 202, re_job.text
    newest = max(int(f.name[:4]) for f in (bot.root / "queue").glob("[0-9][0-9][0-9][0-9].json"))
    rv = bot.job(newest)
    assert rv["kind"] == "voice" and rv["root"] == job_id and [i["voice"] for i in rv["items"]] == [ES_VOICE]


def test_russian_original_targets_english(env):
    client, bot, _ = env
    script = _upload(client, RU, name="Отступление_RU.docx")
    assert script["source_lang"] == "RU" and not script["script_id"].endswith("_RU")
    created = client.post("/api/voice/jobs", json={"script_id": script["script_id"], "mode": "translate",
                                                   "source_lang": "RU", "targets": ["EN", "RU", "ES"],
                                                   "model": "gemini-4-argon-medium"}).json()
    assert created["source_lang"] == "RU" and created["targets"] == ["EN", "ES"]
    assert bot.job(created["id"])["langs"] == ["EN", "ES"]
    bot.translate(script["script_id"], "EN", ["In the winter of 1812 the army retreated along a ruined road.", "The general knew."])
    paras = client.get(f"/api/voice/jobs/{created['id']}/languages/EN/paragraphs").json()
    assert paras[0]["norm"][0] > 100  # English from Russian is longer in characters
    src = client.get(f"/api/voice/jobs/{created['id']}/languages/RU/paragraphs").json()
    assert src[0]["text"] == RU[0] and src[0]["length_pct"] == 100


def test_stop_resume_delete_and_errors(env):
    client, bot, svc = env
    script = _upload(client, EN)
    job_id = client.post("/api/voice/jobs", json={"script_id": script["script_id"], "mode": "translate",
                                                  "targets": ["RU"]}).json()["id"]
    assert client.post(f"/api/voice/jobs/{job_id}/stop").json()["status"] == "stopped"
    assert client.post(f"/api/voice/jobs/{job_id}/resume").json()["status"] == "queued"
    bot.update(job_id, status="running")
    assert client.post(f"/api/voice/jobs/{job_id}/stop").status_code == 200
    assert json.loads((bot.root / "queue" / f"{job_id:04d}.cmd").read_text())["cmd"] == "stop"  # the bot will stop it
    bot.update(job_id, status="failed", error="line 1\nagy не записал перевод после 6 попыток")
    failed = client.get(f"/api/voice/jobs/{job_id}").json()
    assert failed["status"] == "failed" and "agy не записал" in failed["error_message"]
    assert client.delete(f"/api/voice/jobs/{job_id}").status_code == 204
    assert client.get(f"/api/voice/jobs/{job_id}").status_code == 404
    assert client.get("/api/voice/jobs/999").status_code == 404
    bad = client.post("/api/voice/jobs", json={"script_id": "nope", "mode": "all"})
    assert bad.status_code == 404
    unknown = client.post("/api/voice/jobs", json={"script_id": script["script_id"], "model": "gpt-9"})
    assert unknown.status_code == 422


def test_translator_missing_gives_clear_error(tmp_path):
    empty = tmp_path / "nothing"
    empty.mkdir()
    with TestClient(create_app(_settings(tmp_path, empty))) as client:
        r = client.post("/api/voice/scripts", files={"file": ("a.docx", _docx(EN), DOCX)})
        assert r.status_code == 503 and r.json()["error"]["code"] == "translator_missing"
        assert client.get("/api/voice/jobs").json() == []


WORKERS = r'''
import json, subprocess, sys
sys.path[:0] = ["tools", "."]
import imageio_ffmpeg, jobs, voice

def tone(sec):
    return subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), "-v", "error", "-f", "lavfi", "-i",
                           f"sine=frequency=300:duration={sec}", "-ac", "1", "-f", "mp3", "-"], capture_output=True).stdout

voice.synthesize = lambda text, vid, key, model=None, attempts=6: tone(1.0 + len(text) / 300)
voice.missing_voices = lambda items, key=None: []
say = lambda job, text, files=None: None
for _ in range(4):  # translate → voice → mix, the way the bot's workers pick them up
    for engine in ("agy", "claude"):
        jobs.TranslatePool(say, engine).run_once()
    jobs.VoiceWorker(say).loop(once=True)
    jobs.MixWorker(say).loop(once=True)
'''

FAKE_AGY = r'''#!/usr/bin/env python3
import re, sys
from pathlib import Path
prompt = sys.argv[sys.argv.index("--print") + 1]
task = Path(re.search(r"Прочитай файл (.+?) и выполни", prompt).group(1)).read_text(encoding="utf-8")
target = re.search(r"запиши результат в файл (\S+) в таком формате", task).group(1)
lines = {"английский": "The army retreated along the ruined road through the long and bitter winter of that year.",
         "испанский": "El ejército se retiró por el camino arruinado durante aquel largo y amargo invierno."}
lang = re.search(r"Язык перевода: ([^.]+)\.", task).group(1)
keys = re.findall(r"^## (H?\d+) \[", task, re.M)
Path(target).write_text("".join(f"## {k}\n{'Intro' if k.startswith('H') else lines.get(lang, 'текст')}\n" for k in keys),
                        encoding="utf-8")
'''


@pytest.mark.skipif(not os.environ.get("FL_TRANSLATOR_REPO"), reason="needs the real translator (FL_TRANSLATOR_REPO)")
def test_end_to_end_with_real_translator_workers(env):
    """Russian original, everything on: web job → translator workers (agy and Fish faked) → web job done."""
    client, bot, svc = env
    py = os.environ.get("FL_TRANSLATOR_PYTHON") or sys.executable  # needs the translator's requirements
    fake = bot.root / "fake_agy"
    fake.write_text(FAKE_AGY, encoding="utf-8")
    fake.chmod(0o755)
    (bot.root / "tools" / "config.local.json").write_text(json.dumps({"agy": {"binary_path": str(fake)},
                                                                       "usage_tracking": False}), encoding="utf-8")
    (bot.root / "bot_config.json").write_text(json.dumps({"fish_api_key": "x"}), encoding="utf-8")
    script = _upload(client, RU, name="Отступление.docx")
    job_id = client.post("/api/voice/jobs", json={"script_id": script["script_id"], "mode": "all", "source_lang": "RU",
                                                  "targets": ["EN", "ES"], "voices": {"EN": "v1", "ES": "v3"},
                                                  "model": "gemini-4-argon-medium", "voice_model": "drama-3",
                                                  "carry_mood": True}).json()["id"]
    svc.scan()
    r = subprocess.run([py, "-c", WORKERS], cwd=bot.root, capture_output=True, text=True, encoding="utf-8",
                       env=dict(os.environ, PYTHONUTF8="1"), timeout=300)
    assert r.returncode == 0, r.stderr[-2000:]
    svc.scan()
    job = client.get(f"/api/voice/jobs/{job_id}").json()
    assert job["status"] == "done", json.dumps(job, ensure_ascii=False)[:1500]
    assert job["source_lang"] == "RU" and job["targets"] == ["EN", "ES"]
    assert job["voice_model"] == "drama-3" and job["model"] == "gemini-4-argon-medium"
    states = {lang["lang"]: (lang["translate"] or {}).get("state", "-") + "/" + lang["voice"]["state"] + "/" +
              (lang["mix"] or {}).get("state", "-") for lang in job["languages"]}
    assert states == {"RU": "-/done/-", "EN": "done/done/done", "ES": "done/done/done"}
    kinds = {(f["lang"], f["kind"]) for f in job["files"]}
    assert {("EN", "docx"), ("ES", "docx"), ("EN", "mixed_zip"), ("RU", "mixed_zip"), ("RU", "raw_zip")} <= kinds
    chain = [bot.job(n) for n in range(job_id, job_id + 3)]
    assert [j["kind"] for j in chain] == ["translate", "voice", "mix"]
    assert {i["voice"] for i in chain[1]["items"]} == {RU_VOICE, EN_VOICE, ES_VOICE}
    assert any(e["event"] == "job.done" and e["data"]["job_id"] == job_id for e in svc._event_history)
    paras = client.get(f"/api/voice/jobs/{job_id}/languages/EN/paragraphs").json()
    assert paras[0]["audio"]["slot_sec"] and paras[0]["norm"][0] > 100
